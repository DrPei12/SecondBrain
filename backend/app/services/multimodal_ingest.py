"""
Lightweight multimodal ingestion for SecondBrain RAG.

The design follows the same practical boundary as RAG-Anything: parse a source
file into modality-aware content units, then represent those units as grounded
text that the existing vector index can retrieve.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import uuid
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from fastapi import UploadFile

from app.core.config import settings


TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".rst",
    ".log",
    ".json",
    ".yaml",
    ".yml",
    ".html",
    ".htm",
    ".py",
    ".js",
    ".jsx",
    ".ts",
    ".tsx",
    ".css",
    ".scss",
    ".xml",
    ".tex",
}
TABLE_EXTENSIONS = {".csv", ".tsv", ".xlsx", ".xlsm"}
DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".pptx"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff", ".svg"}
AUDIO_EXTENSIONS = {".mp3", ".wav", ".m4a", ".aac", ".ogg", ".flac"}
VIDEO_EXTENSIONS = {".mp4", ".mov", ".mkv", ".webm", ".avi"}
EQUATION_RE = re.compile(
    r"(\$\$.*?\$\$|\$[^$\n]{3,}\$|\\\[.*?\\\]|\\begin\{equation\}.*?\\end\{equation\})",
    re.DOTALL,
)


class _HTMLTextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        text = data.strip()
        if text:
            self.parts.append(text)

    def text(self) -> str:
        return "\n".join(self.parts)


@dataclass
class ExtractedDocument:
    original_filename: str
    stored_name: str
    content_type: str
    size_bytes: int
    sha256: str
    title: str
    modalities: list[str]
    content: str
    warnings: list[str] = field(default_factory=list)

    @property
    def source_url(self) -> str:
        return f"upload://{self.stored_name}"

    @property
    def tags(self) -> list[str]:
        return ["ingested-file", *[f"modality:{modality}" for modality in self.modalities]]

    def response_metadata(self) -> dict[str, Any]:
        return {
            "original_filename": self.original_filename,
            "stored_name": self.stored_name,
            "content_type": self.content_type,
            "size_bytes": self.size_bytes,
            "sha256": self.sha256,
            "modalities": self.modalities,
            "source_url": self.source_url,
            "warnings": self.warnings,
        }


class MultimodalIngestionService:
    async def ingest_upload(
        self,
        upload: UploadFile,
        *,
        title: str = "",
        caption: str = "",
    ) -> ExtractedDocument:
        raw = await upload.read()
        if not raw:
            raise ValueError("Uploaded file is empty")
        if len(raw) > settings.RAG_MAX_UPLOAD_BYTES:
            raise ValueError(
                f"Uploaded file exceeds RAG_MAX_UPLOAD_BYTES={settings.RAG_MAX_UPLOAD_BYTES}"
            )

        original_name = Path(upload.filename or "upload.bin").name
        suffix = Path(original_name).suffix.lower()
        stored_name = f"{uuid.uuid4().hex}_{_safe_filename(original_name)}"
        upload_dir = Path(settings.RAG_UPLOAD_DIR)
        upload_dir.mkdir(parents=True, exist_ok=True)
        (upload_dir / stored_name).write_bytes(raw)

        content_type = upload.content_type or _guess_content_type(suffix)
        sha256 = hashlib.sha256(raw).hexdigest()
        extracted_text, modalities, warnings = self._extract(raw, suffix, content_type, caption)
        if caption.strip():
            extracted_text = f"User supplied caption:\n{caption.strip()}\n\n{extracted_text}".strip()

        title_value = title.strip() or Path(original_name).stem or original_name
        metadata_block = self._metadata_block(
            title=title_value,
            original_name=original_name,
            stored_name=stored_name,
            content_type=content_type,
            size_bytes=len(raw),
            sha256=sha256,
            modalities=modalities,
            warnings=warnings,
        )
        content = f"{metadata_block}\n\n## Extracted Content\n\n{extracted_text}".strip()
        content = content[: settings.RAG_MAX_EXTRACTED_CHARS]

        return ExtractedDocument(
            original_filename=original_name,
            stored_name=stored_name,
            content_type=content_type,
            size_bytes=len(raw),
            sha256=sha256,
            title=title_value,
            modalities=modalities,
            content=content,
            warnings=warnings,
        )

    def _extract(
        self,
        raw: bytes,
        suffix: str,
        content_type: str,
        caption: str,
    ) -> tuple[str, list[str], list[str]]:
        warnings: list[str] = []
        modalities = _modalities_for_file(suffix, content_type)

        if suffix in TABLE_EXTENSIONS:
            return self._extract_table(raw, suffix, warnings), _with_modality(modalities, "table"), warnings
        if suffix in DOCUMENT_EXTENSIONS:
            return self._extract_document(raw, suffix, warnings), modalities, warnings
        if suffix in IMAGE_EXTENSIONS or content_type.startswith("image/"):
            return self._extract_image(raw, suffix, content_type, caption, warnings), modalities, warnings
        if suffix in AUDIO_EXTENSIONS or content_type.startswith("audio/"):
            warnings.append("Audio transcription is not configured; indexed metadata and caption only.")
            return "Audio file metadata was indexed. Add a caption or transcript for semantic retrieval.", modalities, warnings
        if suffix in VIDEO_EXTENSIONS or content_type.startswith("video/"):
            warnings.append("Video transcription is not configured; indexed metadata and caption only.")
            return "Video file metadata was indexed. Add a caption or transcript for semantic retrieval.", modalities, warnings

        text = _decode_text(raw, warnings)
        if suffix in {".html", ".htm"} or "html" in content_type:
            text = _strip_html(text)
        if suffix == ".json":
            text = _format_json(text, warnings)
        if EQUATION_RE.search(text):
            modalities = _with_modality(modalities, "equation")
        if _looks_like_markdown_table(text):
            modalities = _with_modality(modalities, "table")
        return text, modalities, warnings

    def _extract_table(self, raw: bytes, suffix: str, warnings: list[str]) -> str:
        if suffix in {".xlsx", ".xlsm"}:
            try:
                from openpyxl import load_workbook
            except Exception:
                warnings.append("openpyxl is unavailable; spreadsheet metadata indexed only.")
                return "Spreadsheet file metadata was indexed. Install openpyxl to extract worksheets."

            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            sections: list[str] = []
            for sheet in workbook.worksheets[:10]:
                rows = []
                for row in sheet.iter_rows(max_row=80, max_col=20, values_only=True):
                    rows.append(["" if value is None else str(value) for value in row])
                if rows:
                    sections.append(f"### Sheet: {sheet.title}\n\n{_rows_to_markdown(rows)}")
            return "\n\n".join(sections) or "Spreadsheet contained no readable cells."

        delimiter = "\t" if suffix == ".tsv" else ","
        text = _decode_text(raw, warnings)
        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        rows = [row for _, row in zip(range(120), reader)]
        return _rows_to_markdown(rows) if rows else "Table file contained no readable rows."

    def _extract_document(self, raw: bytes, suffix: str, warnings: list[str]) -> str:
        if suffix == ".pdf":
            try:
                from pypdf import PdfReader
            except Exception:
                warnings.append("pypdf is unavailable; PDF metadata indexed only.")
                return "PDF file metadata was indexed. Install pypdf to extract page text."

            reader = PdfReader(io.BytesIO(raw))
            pages = []
            for index, page in enumerate(reader.pages[:80], 1):
                text = page.extract_text() or ""
                if text.strip():
                    pages.append(f"### Page {index}\n\n{text.strip()}")
            return "\n\n".join(pages) or "PDF had no extractable text layer."

        if suffix == ".docx":
            try:
                from docx import Document
            except Exception:
                warnings.append("python-docx is unavailable; DOCX metadata indexed only.")
                return "DOCX metadata was indexed. Install python-docx to extract document text."

            document = Document(io.BytesIO(raw))
            parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table_index, table in enumerate(document.tables, 1):
                rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
                if rows:
                    parts.append(f"### Table {table_index}\n\n{_rows_to_markdown(rows)}")
            return "\n\n".join(parts) or "DOCX contained no readable paragraphs."

        if suffix == ".pptx":
            try:
                from pptx import Presentation
            except Exception:
                warnings.append("python-pptx is unavailable; PPTX metadata indexed only.")
                return "PPTX metadata was indexed. Install python-pptx to extract slide text."

            presentation = Presentation(io.BytesIO(raw))
            slides = []
            for index, slide in enumerate(presentation.slides, 1):
                texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        texts.append(shape.text.strip())
                if texts:
                    slides.append(f"### Slide {index}\n\n" + "\n\n".join(texts))
            return "\n\n".join(slides) or "PPTX contained no readable slide text."

        return "Document metadata was indexed."

    def _extract_image(
        self,
        raw: bytes,
        suffix: str,
        content_type: str,
        caption: str,
        warnings: list[str],
    ) -> str:
        if suffix == ".svg":
            text = _strip_html(_decode_text(raw, warnings))
            return f"SVG image text/markup summary:\n{text[:8000]}"

        details = [f"Image content type: {content_type}"]
        try:
            from PIL import Image

            with Image.open(io.BytesIO(raw)) as image:
                details.append(f"Dimensions: {image.width} x {image.height}")
                details.append(f"Mode: {image.mode}")
                if image.get_format_mimetype():
                    details.append(f"Detected MIME: {image.get_format_mimetype()}")
        except Exception as exc:
            warnings.append(f"Image metadata extraction failed: {exc}")

        if not caption.strip():
            warnings.append(
                "No image caption or OCR/VLM description was provided; semantic retrieval is metadata-only."
            )
        return "\n".join(details)

    @staticmethod
    def _metadata_block(
        *,
        title: str,
        original_name: str,
        stored_name: str,
        content_type: str,
        size_bytes: int,
        sha256: str,
        modalities: list[str],
        warnings: list[str],
    ) -> str:
        warning_text = "\n".join(f"- {warning}" for warning in warnings) or "- none"
        return (
            f"# {title}\n\n"
            "## File Metadata\n\n"
            f"- Original filename: {original_name}\n"
            f"- Stored object: upload://{stored_name}\n"
            f"- Content type: {content_type}\n"
            f"- Size bytes: {size_bytes}\n"
            f"- SHA256: {sha256}\n"
            f"- Modalities: {', '.join(modalities)}\n\n"
            "## Extraction Warnings\n\n"
            f"{warning_text}"
        )


def _safe_filename(filename: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", filename).strip("._")
    return cleaned or "upload.bin"


def _guess_content_type(suffix: str) -> str:
    mapping = {
        ".txt": "text/plain",
        ".md": "text/markdown",
        ".markdown": "text/markdown",
        ".csv": "text/csv",
        ".tsv": "text/tab-separated-values",
        ".json": "application/json",
        ".html": "text/html",
        ".pdf": "application/pdf",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }
    if suffix in IMAGE_EXTENSIONS:
        return f"image/{suffix.lstrip('.')}"
    if suffix in AUDIO_EXTENSIONS:
        return f"audio/{suffix.lstrip('.')}"
    if suffix in VIDEO_EXTENSIONS:
        return f"video/{suffix.lstrip('.')}"
    return mapping.get(suffix, "application/octet-stream")


def _modalities_for_file(suffix: str, content_type: str) -> list[str]:
    if suffix in TABLE_EXTENSIONS:
        return ["table"]
    if suffix in DOCUMENT_EXTENSIONS:
        return ["document", "text"]
    if suffix in IMAGE_EXTENSIONS or content_type.startswith("image/"):
        return ["image"]
    if suffix in AUDIO_EXTENSIONS or content_type.startswith("audio/"):
        return ["audio"]
    if suffix in VIDEO_EXTENSIONS or content_type.startswith("video/"):
        return ["video"]
    return ["text"]


def _with_modality(modalities: list[str], modality: str) -> list[str]:
    return sorted({*modalities, modality})


def _decode_text(raw: bytes, warnings: list[str]) -> str:
    for encoding in ("utf-8-sig", "utf-16", "gb18030", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    warnings.append("Text decoding used replacement characters.")
    return raw.decode("utf-8", errors="replace")


def _format_json(text: str, warnings: list[str]) -> str:
    try:
        return json.dumps(json.loads(text), ensure_ascii=False, indent=2)
    except Exception:
        warnings.append("JSON parsing failed; raw text indexed.")
        return text


def _strip_html(text: str) -> str:
    parser = _HTMLTextExtractor()
    parser.feed(text)
    return parser.text() or text


def _rows_to_markdown(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    normalized = [(row + [""] * width)[:width] for row in rows]
    header = normalized[0]
    body = normalized[1:]
    lines = [
        "| " + " | ".join(_clean_cell(cell) for cell in header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    for row in body:
        lines.append("| " + " | ".join(_clean_cell(cell) for cell in row) + " |")
    return "\n".join(lines)


def _clean_cell(value: str) -> str:
    return str(value).replace("\n", " ").replace("|", "\\|").strip()


def _looks_like_markdown_table(text: str) -> bool:
    lines = [line.strip() for line in text.splitlines()]
    for index, line in enumerate(lines[:-1]):
        if "|" in line and re.search(r"\|\s*:?-{3,}:?\s*\|", lines[index + 1]):
            return True
    return False


multimodal_ingestion_service = MultimodalIngestionService()
